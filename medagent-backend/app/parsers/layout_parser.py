"""File-type router producing a single ParsedDocument contract."""

from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document as DocxDocument

from app.parsers.base import ParseProgress
from app.parsers.cleaning import clean_preserving_structure
from app.parsers.models import ParsedDocument, ParsedPage, TextBlock
from app.parsers.pymupdf_parser import PyMuPDFParser


PARSER_VERSION = "layout-router-v2"


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def parse_document(
    file_path: str,
    file_type: str,
    progress_callback: ParseProgress | None = None,
) -> ParsedDocument:
    path = Path(file_path)
    extension = file_type.lower().lstrip(".")
    if extension == "pdf":
        return clean_preserving_structure(PyMuPDFParser().parse(str(path), progress_callback))

    if extension == "docx":
        source = DocxDocument(str(path))
        blocks: list[TextBlock] = []
        for order, paragraph in enumerate(source.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = (paragraph.style.name or "").lower() if paragraph.style else ""
            kind = "heading" if "heading" in style_name or "标题" in style_name else "text"
            blocks.append(TextBlock(text=text, kind=kind, order=order, metadata={"style": style_name}))
        parsed = ParsedDocument(
            str(path), "docx", PARSER_VERSION,
            [ParsedPage(1, "\n".join(x.text for x in blocks), "docx", blocks=blocks)],
            metadata={"content_sha256": _digest(path)},
        )
        return clean_preserving_structure(parsed)

    if extension in {"txt", "md", "markdown"}:
        raw = path.read_text(encoding="utf-8", errors="replace")
        blocks = []
        for order, line in enumerate(raw.splitlines()):
            text = line.strip()
            if not text:
                continue
            kind = "heading" if text.startswith("#") else "text"
            blocks.append(TextBlock(text=text, kind=kind, order=order))
        parsed = ParsedDocument(
            str(path), extension, PARSER_VERSION,
            [ParsedPage(1, raw, "text", blocks=blocks)],
            metadata={"content_sha256": _digest(path)},
        )
        return clean_preserving_structure(parsed)
    raise ValueError(f"Unsupported file type: {file_type}")
